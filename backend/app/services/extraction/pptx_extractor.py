import io
from typing import List
from pptx import Presentation
from app.domain.documents.models import DocumentPage, ExtractedSection, ExtractedTable
import logging

logger = logging.getLogger(__name__)

class PPTXExtractor:
    async def extract(self, file_bytes: bytes) -> List[DocumentPage]:
        try:
            stream = io.BytesIO(file_bytes)
            prs = Presentation(stream)
        except Exception as e:
            logger.warning(f"Failed to parse pptx presentation ({e}), attempting text fallback.")
            try:
                decoded = file_bytes.decode("utf-8", errors="replace").strip()
                if decoded:
                    return [DocumentPage(page_number=1, text=decoded, sections=[], tables=[])]
            except Exception:
                pass
            raise
        pages: List[DocumentPage] = []

        for idx, slide in enumerate(prs.slides):
            slide_number = idx + 1
            slide_title = f"Slide {slide_number}"
            slide_text_lines: List[str] = []
            tables: List[ExtractedTable] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        slide_text_lines.append(text)
                if shape.has_table:
                    tbl = shape.table
                    rows_data = []
                    for row in tbl.rows:
                        rows_data.append([cell.text.strip() for cell in row.cells])
                    if rows_data:
                        tables.append(ExtractedTable(headers=rows_data[0], rows=rows_data[1:] if len(rows_data) > 1 else []))

            # If slide title shape exists, capture it
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()

            content = "\n".join(slide_text_lines)
            for t in tables:
                md = t.to_markdown()
                if md:
                    content += f"\n\n[Table]\n{md}"

            section = ExtractedSection(title=slide_title, content=content)
            pages.append(
                DocumentPage(
                    page_number=slide_number,
                    text=content,
                    sections=[section],
                    tables=tables
                )
            )

        return pages

pptx_extractor = PPTXExtractor()
