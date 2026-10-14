import io
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

from app.domain.documents.models import DocumentContent, DocumentPage, ExtractedSection
from app.services.extraction.chunking_service import chunking_service
from app.services.extraction.txt_extractor import txt_extractor


@pytest.mark.asyncio
async def test_txt_extractor():
    sample_text = b"Photosynthesis is the process by which green plants use sunlight to synthesize nutrients."
    pages = await txt_extractor.extract(sample_text)
    assert len(pages) == 1
    assert "Photosynthesis" in pages[0].text
    assert pages[0].page_number == 1

@pytest.mark.asyncio
async def test_chunking_preserves_metadata():
    pages = [
        DocumentPage(
            page_number=1,
            text="Cellular respiration converts glucose into ATP.",
            sections=[
                ExtractedSection(
                    title="Glycolysis",
                    content="Glycolysis occurs in the cytoplasm and breaks down glucose."
                )
            ]
        )
    ]
    doc = DocumentContent(
        document_id="doc-test-1",
        title="Biology.txt",
        source_type="txt",
        pages=pages,
        total_pages=1
    )

    chunks = chunking_service.chunk_document(doc)
    assert len(chunks) >= 1
    assert chunks[0].document_id == "doc-test-1"
    assert chunks[0].section == "Glycolysis"
    assert chunks[0].page_start == 1
    assert "Glycolysis occurs in the cytoplasm" in chunks[0].content

@pytest.mark.asyncio
async def test_suggested_topics_extraction():
    from app.services.extraction.extractor_service import extractor_service
    pages = [
        DocumentPage(
            page_number=1,
            text="Cardiovascular anatomy overview.",
            sections=[
                ExtractedSection(title="1.1 Cardiac Muscle Physiology", content="Heart tissue details."),
                ExtractedSection(title="1.2 Coronary Circulation", content="Blood vessel details.")
            ]
        ),
        DocumentPage(
            page_number=2,
            text="Electrical conduction.",
            sections=[
                ExtractedSection(title="1.3 Sinoatrial Node Function", content="Pacemaker details.")
            ]
        )
    ]
    topics = extractor_service.extract_suggested_topics(pages, "Human_Cardiovascular_System.pdf")
    assert len(topics) >= 2
    assert "Entire Document (Human Cardiovascular System)" in topics[0]
    assert any("Cardiac Muscle Physiology" in t for t in topics)
    assert any("Coronary Circulation" in t for t in topics)


def make_docx(paragraphs=(), tables=()):
    import docx

    document = docx.Document()
    for style, text in paragraphs:
        if style:
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    for cells in tables:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = cells
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_docx_pseudo_pages_preserve_heading_and_split_section():
    from app.services.extraction.docx_extractor import docx_extractor

    text = " ".join(["mitochondria"] * 501)
    pages = await docx_extractor.extract(make_docx([(True, "Cell Biology"), (False, text)]))

    assert [page.page_number for page in pages] == [1, 2]
    assert all(page.sections[0].title == "Cell Biology" for page in pages)
    assert sum(len(section.content.split()) for page in pages for section in page.sections) == 503
    assert "Cell Biology" in pages[0].text and "Cell Biology" not in pages[1].text
    assert "mitochondria" in pages[1].text


@pytest.mark.parametrize("word_count, expected_pages", [(25_000, 50), (25_001, None)])
async def test_docx_enforces_fifty_pseudo_page_limit(word_count, expected_pages):
    from app.services.extraction.docx_extractor import docx_extractor

    data = make_docx([(False, " ".join(["biology"] * word_count))])
    if expected_pages is None:
        with pytest.raises(ValueError, match="maximum page count"):
            await docx_extractor.extract(data)
    else:
        pages = await docx_extractor.extract(data)
        assert len(pages) == expected_pages


async def test_docx_table_text_counts_toward_page_limit_and_keeps_table_provenance():
    from app.services.extraction.docx_extractor import docx_extractor

    table_text = " ".join(["enzyme"] * 501)
    pages = await docx_extractor.extract(make_docx([(True, "Enzyme Review")], [table_text]))
    assert len(pages) == 2
    assert pages[0].tables and pages[1].tables == pages[0].tables
    assert all(section.title in {"Enzyme Review", "Table"} for page in pages for section in page.sections)
    assert "enzyme" in pages[1].text
    chunks = chunking_service.chunk_document(
        DocumentContent(document_id="doc-table", title="review.docx", source_type="docx", pages=pages)
    )
    assert any("enzyme" in chunk.content and chunk.page_start == 2 for chunk in chunks)

    over_limit = make_docx([(False, " ".join(["biology"] * 25_000))], ["extra"])
    with pytest.raises(ValueError, match="maximum page count"):
        await docx_extractor.extract(over_limit)


async def test_txt_page_limit_checks_before_creating_page_objects(monkeypatch):
    import app.services.extraction.txt_extractor as module

    pages = await module.txt_extractor.extract(b"x" * (50 * 2500))
    assert len(pages) == 50
    page_constructor = Mock(side_effect=AssertionError("over-limit page was constructed"))
    monkeypatch.setattr(module, "DocumentPage", page_constructor)
    with pytest.raises(ValueError, match="maximum page count"):
        await module.txt_extractor.extract(b"x" * (50 * 2500 + 1))
    page_constructor.assert_not_called()


def make_pptx(slide_count):
    from pptx import Presentation

    presentation = Presentation()
    for _ in range(slide_count):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


async def test_pptx_page_limit_checks_slide_count_before_shape_traversal(monkeypatch):
    import app.services.extraction.pptx_extractor as module
    from app.services.extraction.pptx_extractor import pptx_extractor

    pages = await pptx_extractor.extract(make_pptx(50))
    assert [page.page_number for page in pages] == list(range(1, 51))

    class OverLimitSlides:
        def __len__(self):
            return 51

        def __iter__(self):
            raise AssertionError("over-limit slides were traversed")

    valid_archive = make_pptx(1)
    monkeypatch.setattr(module, "Presentation", lambda _: SimpleNamespace(slides=OverLimitSlides()))
    with pytest.raises(ValueError, match="maximum page count"):
        await pptx_extractor.extract(valid_archive)
